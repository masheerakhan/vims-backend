from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.utils.text import slugify
from rest_framework import status
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.viewsets import ModelViewSet
from rest_framework.views import APIView

from apps.budgets.models import (
    BudgetCategory,
    BudgetSubCategory,
    Budget,
    BudgetLine,
    BudgetRule,
    BudgetConsumption,
    BudgetVarianceRequest,
    BudgetImportBatch,
    BudgetRevision,
    BudgetStatus,
    ImportMode,
    BudgetRevisionSource,
    BudgetRevisionStatus,
)
from apps.budgets.api.serializers import (
    BudgetCategorySerializer,
    BudgetCategoryCreateSerializer,
    BudgetSubCategorySerializer,
    BudgetSubCategoryCreateSerializer,
    BudgetSerializer,
    BudgetCreateSerializer,
    BudgetUpdateSerializer,
    BudgetLineSerializer,
    BudgetLineCreateSerializer,
    BudgetLineUpdateSerializer,
    BudgetRuleSerializer,
    BudgetRuleCreateSerializer,
    BudgetConsumptionSerializer,
    BudgetVarianceRequestSerializer,
    VarianceReviewSerializer,
    ReserveBudgetLineSerializer,
    ConsumeBudgetLineSerializer,
    ReleaseBudgetLineSerializer,
    BudgetImportBatchSerializer,
    BudgetImportBatchListSerializer,
    BudgetImportUploadSerializer,
    BudgetRevisionSerializer,
    BudgetRevisionManualCreateSerializer,
    BudgetRevisionExcelCreateSerializer,
)
from apps.budgets.services import (
    reserve_budget_line,
    consume_reserved_budget_line,
    release_reserved_budget_line,
    review_variance_request,
    resolve_budget_line_for_allocation,
    BudgetLineNotFoundError,
    BudgetLimitExceeded,
    BudgetNotActiveError,
    can_delete_budget,
    can_delete_budget_line,
    can_delete_budget_category,
    can_delete_budget_subcategory,
    can_decrease_budget_allocated,
    can_decrease_budget_line_allocated,
    get_budget_in_use_summary,
    get_budget_line_in_use_summary,
    parse_budget_import_file,
    create_budget_import_batch,
    validate_budget_import_batch,
    commit_budget_import_batch,
    BudgetRevisionConflictError,
    BudgetRevisionValidationError,
    build_budget_revision_snapshot,
    create_budget_revision,
    parse_scoped_budget_revision_file,
    publish_budget_revision,
    resolve_scoped_budget_revision_rows,
)
from apps.budgets.selectors import get_budgets_overview, get_budget_live_balances
from apps.access.selectors import (
    get_user_actionable_scope_ids,
    get_user_actionable_org_ids,
    get_user_visible_scope_ids,
    get_user_visible_org_ids,
)
from apps.access.models import UserRoleAssignment
from apps.access.services import user_can_act_on_scope_or_ancestors_response
from apps.core.models import ScopeNode


def _budget_permission_response(user, scope_node, action: str, action_label: str):
    from apps.access.services import user_has_permission_including_ancestors

    if not user_has_permission_including_ancestors(user, action, "budget", scope_node):
        return Response(
            {"detail": f"You do not have permission to {action_label} at this scope."},
            status=status.HTTP_403_FORBIDDEN,
        )
    return None


# ---------------------------------------------------------------------------
# Category
# ---------------------------------------------------------------------------

class BudgetCategoryViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetCategorySerializer

    def get_queryset(self):
        visible_org_ids = get_user_visible_org_ids(self.request.user)
        qs = BudgetCategory.objects.select_related("org").filter(org_id__in=visible_org_ids)
        org_id = self.request.query_params.get("org")
        is_active = self.request.query_params.get("is_active")
        if org_id:
            qs = qs.filter(org_id=org_id)
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() in ("true", "1", "yes"))
        return qs.order_by("name")

    def get_serializer_class(self):
        if self.action in ("create", "update", "partial_update"):
            return BudgetCategoryCreateSerializer
        return BudgetCategorySerializer

    def _check_org_actionable(self, request, org_id, action_label):
        from apps.core.models import ScopeNode
        scope_nodes = ScopeNode.objects.filter(org_id=org_id).values_list("id", flat=True)
        actionable_ids = get_user_actionable_scope_ids(request.user)
        if not any(sid in actionable_ids for sid in scope_nodes):
            return Response(
                {"detail": f"You do not have permission to {action_label} in this organisation."},
                status=status.HTTP_403_FORBIDDEN,
            )
        return None

    def create(self, request, *args, **kwargs):
        org_id = request.data.get("org")
        if org_id:
            if err := self._check_org_actionable(request, org_id, "create a category"):
                return err
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        category = self.get_object()
        if err := self._check_org_actionable(request, category.org_id, "update this category"):
            return err
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        category = self.get_object()
        if err := self._check_org_actionable(request, category.org_id, "update this category"):
            return err
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        category = self.get_object()
        if err := self._check_org_actionable(request, category.org_id, "delete this category"):
            return err
        ok, reason = can_delete_budget_category(category)
        if not ok:
            return Response({"detail": reason}, status=status.HTTP_409_CONFLICT)
        return super().destroy(request, *args, **kwargs)


# ---------------------------------------------------------------------------
# SubCategory
# ---------------------------------------------------------------------------

class BudgetSubCategoryViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetSubCategorySerializer

    def get_queryset(self):
        visible_org_ids = get_user_visible_org_ids(self.request.user)
        qs = BudgetSubCategory.objects.select_related("category", "category__org").filter(
            category__org_id__in=visible_org_ids
        )
        category_id = self.request.query_params.get("category")
        is_active = self.request.query_params.get("is_active")
        if category_id:
            qs = qs.filter(category_id=category_id)
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() in ("true", "1", "yes"))
        return qs.order_by("category__name", "name")

    def get_serializer_class(self):
        if self.action in ("create", "update", "partial_update"):
            return BudgetSubCategoryCreateSerializer
        return BudgetSubCategorySerializer

    def _check_org_actionable(self, request, org_id, action_label):
        from apps.core.models import ScopeNode
        scope_nodes = ScopeNode.objects.filter(org_id=org_id).values_list("id", flat=True)
        actionable_ids = get_user_actionable_scope_ids(request.user)
        if not any(sid in actionable_ids for sid in scope_nodes):
            return Response(
                {"detail": f"You do not have permission to {action_label} in this organisation."},
                status=status.HTTP_403_FORBIDDEN,
            )
        return None

    def create(self, request, *args, **kwargs):
        category_id = request.data.get("category")
        if category_id:
            try:
                cat = BudgetCategory.objects.get(pk=category_id)
                if err := self._check_org_actionable(request, cat.org_id, "create a subcategory"):
                    return err
            except BudgetCategory.DoesNotExist:
                pass
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        subcategory = self.get_object()
        if err := self._check_org_actionable(request, subcategory.category.org_id, "update this subcategory"):
            return err
        serializer = BudgetSubCategoryCreateSerializer(subcategory, data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response(BudgetSubCategorySerializer(subcategory).data)

    def partial_update(self, request, *args, **kwargs):
        subcategory = self.get_object()
        if err := self._check_org_actionable(request, subcategory.category.org_id, "update this subcategory"):
            return err
        serializer = BudgetSubCategoryCreateSerializer(subcategory, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)
        return Response(BudgetSubCategorySerializer(subcategory).data)

    def destroy(self, request, *args, **kwargs):
        subcategory = self.get_object()
        if err := self._check_org_actionable(request, subcategory.category.org_id, "delete this subcategory"):
            return err
        ok, reason = can_delete_budget_subcategory(subcategory)
        if not ok:
            return Response({"detail": reason}, status=status.HTTP_409_CONFLICT)
        return super().destroy(request, *args, **kwargs)


# ---------------------------------------------------------------------------
# Budget
# ---------------------------------------------------------------------------

class BudgetViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetSerializer

    def get_queryset(self):
        visible_scope_ids = get_user_visible_scope_ids(self.request.user)
        qs = Budget.objects.select_related(
            "org", "scope_node",
        ).prefetch_related("rule", "lines__category", "lines__subcategory").filter(
            scope_node_id__in=visible_scope_ids
        )
        qs = qs.order_by("-created_at")

        for filter_field in ("org", "status"):
            val = self.request.query_params.get(filter_field)
            if val:
                qs = qs.filter(**{filter_field: val})

        scope_node_id = self.request.query_params.get("scope_node")
        if scope_node_id:
            qs = qs.filter(scope_node_id=scope_node_id)

        fy = self.request.query_params.get("financial_year")
        if fy:
            qs = qs.filter(financial_year=fy)

        return qs

    def get_serializer_class(self):
        if self.action == "create":
            return BudgetCreateSerializer
        if self.action in ("update", "partial_update"):
            return BudgetUpdateSerializer
        return BudgetSerializer

    @transaction.atomic
    def create(self, request, *args, **kwargs):
        scope_node_id = request.data.get("scope_node")
        if scope_node_id:
            if err := user_can_act_on_scope_or_ancestors_response(request.user, scope_node_id, "create a budget"):
                return err
            try:
                scope_node = ScopeNode.objects.get(pk=scope_node_id)
            except ScopeNode.DoesNotExist:
                scope_node = None
            if scope_node and (err := _budget_permission_response(request.user, scope_node, "create", "create a budget")):
                return err

        serializer = BudgetCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        lines_data = data.pop("lines", [])

        budget = Budget.objects.create(
            created_by=request.user,
            **data,
        )

        for line_data in lines_data:
            BudgetLine.objects.create(
                budget=budget,
                category=line_data["category"],
                subcategory=line_data.get("subcategory"),
                allocated_amount=line_data["allocated_amount"],
            )

        output = BudgetSerializer(budget, context={"request": request})
        return Response(output.data, status=status.HTTP_201_CREATED)

    @transaction.atomic
    def update(self, request, *args, **kwargs):
        budget = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, budget.scope_node_id, "update this budget"):
            return err
        if err := _budget_permission_response(request.user, budget.scope_node, "update", "update this budget"):
            return err

        partial = kwargs.pop("partial", False)
        serializer = BudgetUpdateSerializer(instance=budget, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        lines_data = data.pop("lines", None)

        # Guard against decreasing allocated below reserved+consumed
        new_allocated = data.get("allocated_amount")
        if new_allocated is not None:
            ok, reason = can_decrease_budget_allocated(budget, new_allocated)
            if not ok:
                return Response({"detail": reason}, status=status.HTTP_409_CONFLICT)

        # Update header fields
        for field, value in data.items():
            setattr(budget, field, value)
        budget.save()

        # Nested line upsert
        if lines_data is not None:
            payload_line_ids = {l["id"] for l in lines_data if "id" in l}

            # Delete lines omitted from payload (only if zero usage)
            for existing_line in budget.lines.all():
                if existing_line.id not in payload_line_ids:
                    ok, reason = can_delete_budget_line(existing_line)
                    if not ok:
                        raise ValidationError(
                            f"Cannot remove line {existing_line.id}: {reason}"
                        )
                    existing_line.delete()

            # Upsert lines from payload
            for line_data in lines_data:
                line_id = line_data.pop("id", None)
                if line_id:
                    line = BudgetLine.objects.get(pk=line_id, budget=budget)
                    incoming_category = line_data.get("category", line.category)
                    incoming_subcategory = line_data.get("subcategory", line.subcategory)
                    category_changed = incoming_category.id != line.category_id
                    subcategory_changed = (
                        (incoming_subcategory.id if incoming_subcategory else None)
                        != line.subcategory_id
                    )
                    if category_changed or subcategory_changed:
                        line_summary = get_budget_line_in_use_summary(line)
                        if line_summary["is_in_use"]:
                            raise ValidationError(
                                f"Cannot change category/subcategory on line {line_id}: "
                                "line has operational history or active usage."
                            )
                    if (category_changed or subcategory_changed) and (
                        incoming_subcategory and incoming_subcategory.category_id != incoming_category.id
                    ):
                        raise ValidationError(
                            f"Cannot change category/subcategory on line {line_id}: "
                            "subcategory does not belong to selected category."
                        )
                    # Guard line allocated decrease
                    if (
                        "allocated_amount" in line_data
                        and line_data["allocated_amount"] != line.allocated_amount
                    ):
                        ok, reason = can_decrease_budget_line_allocated(line, line_data["allocated_amount"])
                        if not ok:
                            raise ValidationError(f"Line {line_id}: {reason}")
                    for field, value in line_data.items():
                        setattr(line, field, value)
                    line.save()
                else:
                    BudgetLine.objects.create(
                        budget=budget,
                        category=line_data["category"],
                        subcategory=line_data.get("subcategory"),
                        allocated_amount=line_data["allocated_amount"],
                    )

        output = BudgetSerializer(budget, context={"request": request})
        return Response(output.data)

    def partial_update(self, request, *args, **kwargs):
        kwargs["partial"] = True
        return self.update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        budget = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, budget.scope_node_id, "delete this budget"):
            return err
        if err := _budget_permission_response(request.user, budget.scope_node, "delete", "delete this budget"):
            return err
        ok, reason = can_delete_budget(budget)
        if not ok:
            return Response({"detail": reason}, status=status.HTTP_409_CONFLICT)
        return super().destroy(request, *args, **kwargs)

    def _resolve_line(self, budget: Budget, budget_line_id: int):
        try:
            return BudgetLine.objects.get(pk=budget_line_id, budget=budget)
        except BudgetLine.DoesNotExist:
            return None

    @action(detail=True, methods=["get"], url_path="live-balances")
    def live_balances(self, request, pk=None):
        """GET /budgets/{id}/live-balances/ — real-time ledger balances."""
        budget = self.get_object()
        return Response(get_budget_live_balances(budget))

    @action(detail=True, methods=["get"], url_path="in-use")
    def in_use(self, request, pk=None):
        """GET /budgets/{id}/in-use/ — whether this budget has active usage."""
        budget = self.get_object()
        return Response(get_budget_in_use_summary(budget))

    @action(detail=True, methods=["get"], url_path="revision-template")
    def revision_template(self, request, pk=None):
        """Download a scoped allocation workbook for this budget only."""
        budget = self.get_object()
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return Response(
                {"detail": "openpyxl is required to generate budget templates."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Budget Allocation"
        headers = [
            "Category Code *",
            "Category Name",
            "Subcategory Code",
            "Subcategory Name",
            "Current Allocation",
            "New Allocation *",
        ]
        worksheet.append(headers)
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="EA580C")

        for line in budget.lines.select_related("category", "subcategory").order_by(
            "category__code", "subcategory__code", "id"
        ):
            worksheet.append([
                line.category.code,
                line.category.name,
                line.subcategory.code if line.subcategory_id else "",
                line.subcategory.name if line.subcategory_id else "",
                float(line.allocated_amount),
                float(line.allocated_amount),
            ])

        worksheet.freeze_panes = "A2"
        for column, width in {
            "A": 28, "B": 38, "C": 32, "D": 42, "E": 20, "F": 20,
        }.items():
            worksheet.column_dimensions[column].width = width

        from io import BytesIO
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f"{slugify(budget.code) or 'budget'}-allocation-template.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    @action(detail=True, methods=["post"], url_path="reserve")
    def reserve(self, request, pk=None):
        """POST /budgets/{id}/reserve/ — reserve amount against a budget line."""
        budget = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, budget.scope_node_id, "reserve budget"):
            return err

        serializer = ReserveBudgetLineSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        line = self._resolve_line(budget, data["budget_line_id"])
        if line is None:
            return Response(
                {"detail": f"Budget line {data['budget_line_id']} not found on this budget."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            result = reserve_budget_line(
                line=line,
                amount=data["amount"],
                source_type=data["source_type"],
                source_id=str(data["source_id"]),
                requested_by=request.user,
                note=data.get("note", ""),
            )
        except BudgetNotActiveError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except BudgetLimitExceeded as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            "status": result["status"],
            "projected_utilization": str(result["projected_utilization"]),
            "current_utilization": str(result["current_utilization"]),
            "consumption": (
                BudgetConsumptionSerializer(result["consumption"]).data
                if result["consumption"]
                else None
            ),
            "variance_request": (
                BudgetVarianceRequestSerializer(result["variance_request"]).data
                if result["variance_request"]
                else None
            ),
        }, status=status.HTTP_201_CREATED if result["consumption"] else status.HTTP_200_OK)

    @action(detail=True, methods=["post"], url_path="consume")
    def consume(self, request, pk=None):
        """POST /budgets/{id}/consume/ — consume from reserved amount on a budget line."""
        budget = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, budget.scope_node_id, "consume budget"):
            return err

        serializer = ConsumeBudgetLineSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        line = self._resolve_line(budget, data["budget_line_id"])
        if line is None:
            return Response(
                {"detail": f"Budget line {data['budget_line_id']} not found on this budget."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            result = consume_reserved_budget_line(
                line=line,
                amount=data["amount"],
                source_type=data["source_type"],
                source_id=str(data["source_id"]),
                consumed_by=request.user,
                note=data.get("note", ""),
            )
        except (BudgetNotActiveError, ValueError) as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            "status": result["status"],
            "consumption": BudgetConsumptionSerializer(result["consumption"]).data,
        }, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"], url_path="release")
    def release(self, request, pk=None):
        """POST /budgets/{id}/release/ — release reserved amount on a budget line."""
        budget = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, budget.scope_node_id, "release budget"):
            return err

        serializer = ReleaseBudgetLineSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        line = self._resolve_line(budget, data["budget_line_id"])
        if line is None:
            return Response(
                {"detail": f"Budget line {data['budget_line_id']} not found on this budget."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            result = release_reserved_budget_line(
                line=line,
                amount=data["amount"],
                source_type=data["source_type"],
                source_id=str(data["source_id"]),
                released_by=request.user,
                note=data.get("note", ""),
            )
        except (BudgetNotActiveError, ValueError) as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response({
            "status": result["status"],
            "consumption": BudgetConsumptionSerializer(result["consumption"]).data,
        }, status=status.HTTP_201_CREATED)


# ---------------------------------------------------------------------------
# BudgetLine (standalone CRUD)
# ---------------------------------------------------------------------------

class BudgetLineViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetLineSerializer
    http_method_names = ["get", "post", "patch", "delete"]

    def get_queryset(self):
        visible_scope_ids = get_user_visible_scope_ids(self.request.user)
        qs = BudgetLine.objects.select_related(
            "budget", "budget__scope_node", "category", "subcategory"
        ).filter(budget__scope_node_id__in=visible_scope_ids)

        budget_id = self.request.query_params.get("budget")
        if budget_id:
            qs = qs.filter(budget_id=budget_id)
        category_id = self.request.query_params.get("category")
        if category_id:
            qs = qs.filter(category_id=category_id)
        return qs.order_by("budget_id", "category__name")

    def get_serializer_class(self):
        if self.action == "create":
            return BudgetLineCreateSerializer
        if self.action in ("update", "partial_update"):
            return BudgetLineUpdateSerializer
        return BudgetLineSerializer

    def create(self, request, *args, **kwargs):
        budget_id = request.data.get("budget")
        if budget_id:
            try:
                budget = Budget.objects.get(pk=budget_id)
                if err := user_can_act_on_scope_or_ancestors_response(request.user, budget.scope_node_id, "add a budget line"):
                    return err
            except Budget.DoesNotExist:
                pass

        serializer = BudgetLineCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        line = BudgetLine.objects.create(
            budget=data["budget"],
            category=data["category"],
            subcategory=data.get("subcategory"),
            allocated_amount=data["allocated_amount"],
        )
        output = BudgetLineSerializer(line, context={"request": request})
        return Response(output.data, status=status.HTTP_201_CREATED)

    def update(self, request, *args, **kwargs):
        line = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, line.budget.scope_node_id, "update this budget line"):
            return err

        partial = kwargs.pop("partial", False)
        serializer = BudgetLineUpdateSerializer(data=request.data, partial=partial)
        serializer.instance = line
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        # Guard line allocated decrease
        if "allocated_amount" in data:
            ok, reason = can_decrease_budget_line_allocated(line, data["allocated_amount"])
            if not ok:
                return Response({"detail": reason}, status=status.HTTP_409_CONFLICT)

        for field, value in data.items():
            setattr(line, field, value)
        line.save()

        output = BudgetLineSerializer(line, context={"request": request})
        return Response(output.data)

    def partial_update(self, request, *args, **kwargs):
        kwargs["partial"] = True
        return self.update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        line = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, line.budget.scope_node_id, "delete this budget line"):
            return err
        ok, reason = can_delete_budget_line(line)
        if not ok:
            return Response({"detail": reason}, status=status.HTTP_409_CONFLICT)
        return super().destroy(request, *args, **kwargs)


# ---------------------------------------------------------------------------
# BudgetRule
# ---------------------------------------------------------------------------

class BudgetRuleViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetRuleSerializer

    def get_queryset(self):
        visible_scope_ids = get_user_visible_scope_ids(self.request.user)
        qs = BudgetRule.objects.select_related("budget").filter(
            budget__scope_node_id__in=visible_scope_ids
        ).order_by("id")
        budget_id = self.request.query_params.get("budget")
        is_active = self.request.query_params.get("is_active")
        if budget_id:
            qs = qs.filter(budget_id=budget_id)
        if is_active is not None:
            qs = qs.filter(is_active=is_active.lower() in ("true", "1", "yes"))
        return qs

    def get_serializer_class(self):
        if self.action in ("create", "update", "partial_update"):
            return BudgetRuleCreateSerializer
        return BudgetRuleSerializer

    def create(self, request, *args, **kwargs):
        budget_id = request.data.get("budget")
        if budget_id:
            try:
                budget = Budget.objects.get(pk=budget_id)
                if err := user_can_act_on_scope_or_ancestors_response(request.user, budget.scope_node_id, "create a budget rule"):
                    return err
            except Budget.DoesNotExist:
                pass
        return super().create(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        rule = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, rule.budget.scope_node_id, "update this budget rule"):
            return err
        return super().update(request, *args, **kwargs)

    def partial_update(self, request, *args, **kwargs):
        rule = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, rule.budget.scope_node_id, "update this budget rule"):
            return err
        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        rule = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(request.user, rule.budget.scope_node_id, "delete this budget rule"):
            return err
        return super().destroy(request, *args, **kwargs)


# ---------------------------------------------------------------------------
# BudgetConsumption (read-only)
# ---------------------------------------------------------------------------

class BudgetConsumptionViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetConsumptionSerializer
    http_method_names = ["get"]

    def get_queryset(self):
        visible_scope_ids = get_user_visible_scope_ids(self.request.user)
        qs = BudgetConsumption.objects.select_related(
            "budget",
            "budget__scope_node",
            "budget_line",
            "budget_line__category",
            "budget_line__subcategory",
            "created_by",
        ).filter(
            budget__scope_node_id__in=visible_scope_ids
        ).order_by("-created_at")
        budget_id = self.request.query_params.get("budget")
        budget_line_id = self.request.query_params.get("budget_line")
        source_type = self.request.query_params.get("source_type")
        source_id = self.request.query_params.get("source_id")
        category = self.request.query_params.get("category")
        subcategory = self.request.query_params.get("subcategory")
        if budget_id:
            qs = qs.filter(budget_id=budget_id)
        if budget_line_id:
            qs = qs.filter(budget_line_id=budget_line_id)
        if source_type:
            qs = qs.filter(source_type=source_type)
        if source_id:
            qs = qs.filter(source_id=source_id)
        if category:
            cat_ids = [c.strip() for c in category.split(",") if c.strip()]
            if cat_ids:
                qs = qs.filter(budget_line__category_id__in=cat_ids)
        if subcategory:
            sub_ids = [s.strip() for s in subcategory.split(",") if s.strip()]
            if sub_ids:
                qs = qs.filter(budget_line__subcategory_id__in=sub_ids)
        date_from = self.request.query_params.get("date_from")
        date_to = self.request.query_params.get("date_to")
        if date_from:
            qs = qs.filter(created_at__date__gte=date_from)
        if date_to:
            qs = qs.filter(created_at__date__lte=date_to)
        return qs

    def get_serializer(self, *args, **kwargs):
        if kwargs.get("many") and args:
            items = list(args[0])
            args = (items, *args[1:])
            invoice_ids = [
                int(c.source_id)
                for c in items
                if getattr(c, "source_type", None) == "invoice"
                and getattr(c, "source_id", None)
                and str(c.source_id).isdigit()
            ]
            invoice_map = {}
            if invoice_ids:
                from apps.invoices.models import Invoice
                for inv in (
                    Invoice.objects
                    .select_related("vendor", "scope_node")
                    .filter(id__in=set(invoice_ids))
                ):
                    invoice_map[str(inv.id)] = inv
            ctx = dict(kwargs.get("context") or self.get_serializer_context())
            ctx["invoice_map"] = invoice_map
            kwargs["context"] = ctx
        return super().get_serializer(*args, **kwargs)


# ---------------------------------------------------------------------------
# BudgetVarianceRequest
# ---------------------------------------------------------------------------

class BudgetVarianceRequestViewSet(ModelViewSet):
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetVarianceRequestSerializer

    def get_queryset(self):
        visible_scope_ids = get_user_visible_scope_ids(self.request.user)
        qs = BudgetVarianceRequest.objects.select_related(
            "budget", "budget_line", "requested_by", "reviewed_by",
        ).filter(budget__scope_node_id__in=visible_scope_ids).order_by("-created_at")
        budget_id = self.request.query_params.get("budget")
        variance_status = self.request.query_params.get("status")
        if budget_id:
            qs = qs.filter(budget_id=budget_id)
        if variance_status:
            qs = qs.filter(status=variance_status)
        return qs

    def create(self, request, *args, **kwargs):
        return Response(
            {"detail": "Variance requests are created automatically by budget reservations."},
            status=405,
        )

    @action(detail=True, methods=["post"], url_path="review")
    def review(self, request, pk=None):
        """POST /budgets/variance-requests/{id}/review/"""
        variance_req = self.get_object()
        if err := user_can_act_on_scope_or_ancestors_response(
            request.user, variance_req.budget.scope_node_id, "review variance request"
        ):
            return err
        serializer = VarianceReviewSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        try:
            updated = review_variance_request(
                variance_request=variance_req,
                decision=data["decision"],
                reviewed_by=request.user,
                review_note=data.get("review_note", ""),
            )
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(BudgetVarianceRequestSerializer(updated).data)


# ---------------------------------------------------------------------------
# BudgetImportBatch
# ---------------------------------------------------------------------------

class BudgetImportBatchViewSet(ModelViewSet):
    """
    Endpoints for bulk Excel import of budget data.

    POST /import-batches/upload/   — parse + create batch + rows
    POST /import-batches/{id}/validate/  — validate all rows in a batch
    POST /import-batches/{id}/commit/    — commit all valid rows
    GET  /import-batches/          — list batches (no rows)
    GET  /import-batches/{id}/     — detail with rows
    """
    permission_classes = [IsAuthenticated]
    http_method_names = ["get", "post"]

    def _tenant_admin_org_ids(self):
        """Return organizations where the caller has an active Tenant Admin role."""
        if self.request.user.is_superuser:
            return None
        return set(
            UserRoleAssignment.objects.filter(
                user=self.request.user,
                role__code="tenant_admin",
                role__is_active=True,
            ).values_list("role__org_id", flat=True)
        )

    def initial(self, request, *args, **kwargs):
        super().initial(request, *args, **kwargs)
        if not request.user.is_superuser and not self._tenant_admin_org_ids():
            raise PermissionDenied("Only Tenant Admin users can manage bulk budget imports.")

    def get_queryset(self):
        admin_org_ids = self._tenant_admin_org_ids()
        if admin_org_ids is None:
            admin_org_ids = get_user_visible_org_ids(self.request.user)
        return BudgetImportBatch.objects.filter(org_id__in=admin_org_ids).order_by("-created_at")

    def get_serializer_class(self):
        if self.action == "list":
            return BudgetImportBatchListSerializer
        return BudgetImportBatchSerializer

    def create(self, request, *args, **kwargs):
        return Response(
            {"detail": "Use /import-batches/upload/ to upload a file."},
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )

    @action(detail=False, methods=["post"], url_path="upload")
    def upload(self, request):
        """
        POST /import-batches/upload/
        Accepts multipart/form-data with `file` (xlsx), optional `financial_year`,
        and optional `import_mode` (setup_only | safe_update | full).
        Creates BudgetImportBatch + rows. Does not validate yet.

        import_mode defaults to SAFE_UPDATE.
        """
        serializer = BudgetImportUploadSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        uploaded_file = data["file"]
        financial_year = data.get("financial_year", "")
        import_mode = data.get("import_mode", ImportMode.SAFE_UPDATE)

        # Resolve the organization only from the caller's Tenant Admin role scope.
        from apps.core.models import Organization
        admin_org_ids = self._tenant_admin_org_ids()
        if admin_org_ids is None:
            admin_org_ids = set(get_user_visible_org_ids(request.user))
        org_ids = list(admin_org_ids)
        if not org_ids:
            return Response({"detail": "No Tenant Admin organisation assignment found."}, status=status.HTTP_403_FORBIDDEN)

        requested_org_id = request.data.get("org")
        if requested_org_id is not None and str(requested_org_id) not in {str(org_id) for org_id in org_ids}:
            return Response(
                {"detail": "You cannot create a bulk budget import for this organization."},
                status=status.HTTP_403_FORBIDDEN,
            )
        try:
            org = Organization.objects.get(pk=requested_org_id or org_ids[0])
        except Organization.DoesNotExist:
            return Response({"detail": "Organization not found."}, status=status.HTTP_400_BAD_REQUEST)

        try:
            parsed = parse_budget_import_file(uploaded_file)
        except (ValueError, ImportError) as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)

        if not parsed:
            return Response({"detail": "The file has no data rows."}, status=status.HTTP_400_BAD_REQUEST)

        batch = create_budget_import_batch(
            org=org,
            file_name=uploaded_file.name,
            parsed_rows=parsed,
            created_by=request.user,
            financial_year=financial_year,
            import_mode=import_mode,
        )
        return Response(BudgetImportBatchListSerializer(batch).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"], url_path="validate")
    def validate_batch(self, request, pk=None):
        """POST /import-batches/{id}/validate/ — run validation on all pending rows."""
        batch = self.get_object()
        try:
            batch = validate_budget_import_batch(batch)
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response(BudgetImportBatchSerializer(batch).data)

    @action(detail=True, methods=["post"], url_path="commit")
    def commit(self, request, pk=None):
        """POST /import-batches/{id}/commit/ — commit all valid rows to budget/line records."""
        batch = self.get_object()
        try:
            batch = commit_budget_import_batch(batch, committed_by=request.user)
        except ValueError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        return Response(BudgetImportBatchSerializer(batch).data)


# ---------------------------------------------------------------------------
# Scoped budget revisions
# ---------------------------------------------------------------------------

class BudgetRevisionViewSet(ModelViewSet):
    """
    Versioned allocation changes for a single selected budget.

    POST /revisions/manual/ accepts one full allocation plan from the manual UI.
    POST /revisions/excel/ accepts one selected budget's allocation workbook.
    POST /revisions/{id}/publish/ applies the validated revision atomically.
    """
    permission_classes = [IsAuthenticated]
    serializer_class = BudgetRevisionSerializer
    http_method_names = ["get", "post", "head", "options"]

    def get_queryset(self):
        visible_scope_ids = get_user_visible_scope_ids(self.request.user)
        qs = (
            BudgetRevision.objects.select_related(
                "budget", "budget__scope_node", "created_by", "published_by"
            )
            .prefetch_related("lines__category", "lines__subcategory")
            .filter(budget__scope_node_id__in=visible_scope_ids)
            .order_by("-created_at")
        )
        budget_id = self.request.query_params.get("budget")
        if budget_id:
            qs = qs.filter(budget_id=budget_id)
        revision_status = self.request.query_params.get("status")
        if revision_status:
            qs = qs.filter(status=revision_status)
        return qs

    def create(self, request, *args, **kwargs):
        return Response(
            {"detail": "Use /revisions/manual/ or /revisions/excel/ to create a budget revision."},
            status=status.HTTP_405_METHOD_NOT_ALLOWED,
        )

    def _can_manage_budget(self, request, budget):
        if err := user_can_act_on_scope_or_ancestors_response(
            request.user, budget.scope_node_id, "change this budget"
        ):
            return err
        return _budget_permission_response(
            request.user, budget.scope_node, "update", "change this budget"
        )

    @action(detail=False, methods=["post"], url_path="manual")
    def create_manual(self, request):
        serializer = BudgetRevisionManualCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        budget = data["budget"]
        if err := self._can_manage_budget(request, budget):
            return err

        try:
            revision = create_budget_revision(
                budget=budget,
                proposed_lines=data["lines"],
                source=BudgetRevisionSource.MANUAL,
                change_reason=data["change_reason"],
                created_by=request.user,
            )
        except BudgetRevisionValidationError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(BudgetRevisionSerializer(revision, context={"request": request}).data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=["post"], url_path="excel")
    def create_excel(self, request):
        serializer = BudgetRevisionExcelCreateSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        budget = data["budget"]
        if err := self._can_manage_budget(request, budget):
            return err

        uploaded_file = data["file"]
        try:
            parsed_rows = parse_scoped_budget_revision_file(uploaded_file)
            proposed_lines = resolve_scoped_budget_revision_rows(budget, parsed_rows)
            uploaded_file.seek(0)
            revision = create_budget_revision(
                budget=budget,
                proposed_lines=proposed_lines,
                source=BudgetRevisionSource.EXCEL,
                change_reason=data["change_reason"],
                created_by=request.user,
                source_file=uploaded_file,
                source_file_name=uploaded_file.name,
            )
        except (BudgetRevisionValidationError, ImportError) as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)

        return Response(BudgetRevisionSerializer(revision, context={"request": request}).data, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=["post"], url_path="publish")
    def publish(self, request, pk=None):
        revision = self.get_object()
        if err := self._can_manage_budget(request, revision.budget):
            return err
        try:
            revision = publish_budget_revision(revision=revision, published_by=request.user)
        except BudgetRevisionConflictError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_409_CONFLICT)
        except BudgetRevisionValidationError as exc:
            return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        return Response(BudgetRevisionSerializer(revision, context={"request": request}).data)

    @action(detail=True, methods=["post"], url_path="cancel")
    def cancel(self, request, pk=None):
        revision = self.get_object()
        if err := self._can_manage_budget(request, revision.budget):
            return err
        if revision.status not in (BudgetRevisionStatus.DRAFT, BudgetRevisionStatus.VALIDATED):
            return Response(
                {"detail": f"Revision {revision.id} is '{revision.status}' and cannot be cancelled."},
                status=status.HTTP_400_BAD_REQUEST,
            )
        revision.status = BudgetRevisionStatus.CANCELLED
        revision.save(update_fields=["status", "updated_at"])
        return Response(BudgetRevisionSerializer(revision, context={"request": request}).data)


# ---------------------------------------------------------------------------
# Budget Overview (analytics dashboard)
# ---------------------------------------------------------------------------

class BudgetOverviewView(APIView):
    """GET /api/v1/budgets/overview/ — aggregated budget analytics."""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        payload = get_budgets_overview(request.user)
        return Response(payload)
