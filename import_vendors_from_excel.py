"""
Standalone vendor import script for FundFlow V2.

Usage:
    python import_vendors_from_excel.py <path_to_excel_file.xlsx> [--dry-run]

Features:
    - Reads Excel sheet (.xlsx) using openpyxl.
    - Flexible column matching (case-insensitive, trims spaces).
    - Vendors WITH code: marked as ACTIVE & APPROVED (operational_status='active', marketing_status='approved').
    - Vendors WITHOUT code: assigned provisional 'SAP-AUTO-...' ID and marked as ACTIVE & PENDING (operational_status='active', marketing_status='pending').
    - Resolves ScopeNode (Park/Department) if column provided, with graceful fallback.
    - Supports --dry-run to preview what will be created/updated before saving.
"""

import os
import sys
import argparse
from uuid import uuid4
from pathlib import Path

# Setup Django environment
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings.local")
import django
django.setup()

import openpyxl
from django.db import transaction
from django.utils import timezone
from django.utils.text import slugify

from apps.core.models import ScopeNode, Organization
from apps.vendors.models import Vendor, OperationalStatus, MarketingStatus


# Flexible header aliases (lowercase, stripped)
HEADER_ALIASES = {
    "sap_vendor_id": [
        "vendor code", "vendor_code", "sap vendor id", "sap_vendor_id",
        "vendor id", "vendor_id", "sap code", "sap_code", "sap id", "sap_id",
        "sap", "sap no", "sap number", "code",
    ],
    "vendor_name": [
        "vendor name", "vendor_name", "vendor", "name", "company name",
        "company", "vendor / supplier name", "supplier name", "party name",
    ],
    "email": [
        "email", "vendor email", "vendor_email", "email address", "email id",
        "mail", "contact email", "primary email",
    ],
    "phone": [
        "phone", "mobile", "contact", "phone number", "contact number",
        "contact phone", "mobile no", "phone no",
    ],
    "scope": [
        "department / scope code", "scope code", "scope_node_code", "scope",
        "department", "park", "location", "entity", "site", "park name",
    ],
    "pan": ["pan", "pan number", "pan_no", "pan no"],
    "gstin": ["gstin", "gst", "gst number", "gst_no", "gstin no"],
    "address": ["address line 1", "address_line1", "address", "street"],
    "city": ["city"],
    "state": ["state"],
    "pincode": ["pincode", "pin code", "zip", "postal code"],
    "bank_name": ["bank name", "bank_name", "bank"],
    "account_number": ["account number", "account_number", "bank account number", "account no", "a/c no"],
    "ifsc": ["ifsc", "ifsc code", "bank ifsc"],
    "beneficiary_name": ["beneficiary name", "beneficiary_name", "account holder name"],
}


def clean_cell_value(val) -> str:
    """Format cell value cleanly, stripping .0 from integer floats."""
    if val is None:
        return ""
    if isinstance(val, float):
        if val.is_integer():
            return str(int(val)).strip()
        return str(val).strip()
    if isinstance(val, int):
        return str(val).strip()
    return str(val).strip()


def normalize_header(header_val: str) -> str:
    """Normalize raw header string to standard internal key."""
    if not header_val:
        return ""
    clean = str(header_val).strip().lower()
    clean = clean.replace("*", "").strip()

    # Pass 1: EXACT MATCH FIRST across all aliases (prevents prefix collision like 'vendor code' matching 'vendor')
    for standard_key, aliases in HEADER_ALIASES.items():
        if clean in aliases:
            return standard_key

    # Pass 2: Prefix match sorted by longest alias first
    for standard_key, aliases in sorted(HEADER_ALIASES.items(), key=lambda item: max(len(a) for a in item[1]), reverse=True):
        for alias in sorted(aliases, key=len, reverse=True):
            if clean == alias or clean.startswith(alias + " "):
                return standard_key

    return clean


def resolve_scope_node(raw_scope_str: str, default_scope: ScopeNode) -> ScopeNode:
    """Attempt to find ScopeNode by code or name, fallback to default."""
    if not raw_scope_str:
        return default_scope
    
    val = str(raw_scope_str).strip()
    # Try exact code
    node = ScopeNode.objects.filter(code__iexact=val).first()
    if node:
        return node
    # Try exact name
    node = ScopeNode.objects.filter(name__iexact=val).first()
    if node:
        return node
    # Try contains name
    node = ScopeNode.objects.filter(name__icontains=val).first()
    if node:
        return node
    return default_scope


def import_vendors(file_path: str, dry_run: bool = False):
    path = Path(file_path)
    if not path.is_file():
        print(f"[ERROR] File not found: {file_path}")
        sys.exit(1)

    print("=" * 80)
    print(f"LOADING WORKBOOK: {path.name}")
    print(f"MODE: {'DRY RUN (No changes will be saved)' if dry_run else 'LIVE COMMIT (Saving to database)'}")
    print("=" * 80)

    wb = openpyxl.load_workbook(file_path, data_only=True)
    sheet = wb.active

    # Find header row
    header_col_map = {}
    header_row_idx = 1
    for r in range(1, min(10, sheet.max_row + 1)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        mapped = [normalize_header(clean_cell_value(v)) for v in row_vals if v is not None]
        if "vendor_name" in mapped or "sap_vendor_id" in mapped:
            header_row_idx = r
            for col_idx, raw_val in enumerate(row_vals, 1):
                if raw_val is not None:
                    norm = normalize_header(clean_cell_value(raw_val))
                    if norm:
                        header_col_map[col_idx] = norm
            break

    if "vendor_name" not in header_col_map.values():
        print("[ERROR] Could not find 'Vendor Name' column in header row.")
        print(f"Detected columns: {list(header_col_map.values())}")
        sys.exit(1)

    print(f"Detected Headers at Row {header_row_idx}: {header_col_map}")

    # Fallback scope node & organization
    default_org = Organization.objects.first()
    default_scope = ScopeNode.objects.filter(parent__isnull=True).first() or ScopeNode.objects.first()

    if not default_scope:
        print("[ERROR] No ScopeNode exists in database. Please configure core hierarchy first.")
        sys.exit(1)

    total_rows = 0
    created_approved = 0
    created_pending = 0
    updated_count = 0
    skipped_count = 0

    now = timezone.now()

    def process_rows():
        nonlocal total_rows, created_approved, created_pending, updated_count, skipped_count

        for row_idx in range(header_row_idx + 1, sheet.max_row + 1):
            row_data = {}
            for col_idx, field_key in header_col_map.items():
                val = sheet.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    row_data[field_key] = clean_cell_value(val)

            v_name = row_data.get("vendor_name", "").strip()
            if not v_name:
                continue

            total_rows += 1
            raw_code = row_data.get("sap_vendor_id", "").strip()
            v_email = row_data.get("email", "").strip()
            v_phone = row_data.get("phone", "").strip()
            v_pan = row_data.get("pan", "").strip().upper()
            v_gstin = row_data.get("gstin", "").strip().upper()

            # Resolve scope
            scope_str = row_data.get("scope", "")
            scope_node = resolve_scope_node(scope_str, default_scope)
            org = scope_node.org if scope_node else default_org

            has_code = bool(raw_code and raw_code.lower() not in ("none", "null", "n/a", "-", ""))

            if has_code:
                # ── Vendors WITH Code -> Approved & Active ──────────────────
                final_sap_id = raw_code
                op_status = OperationalStatus.ACTIVE
                mkt_status = MarketingStatus.APPROVED
                appr_date = now
                type_label = "ACTIVE & APPROVED (Has Code)"
            else:
                # ── Vendors WITHOUT Code -> Active & Pending ────────────────
                clean_slug = slugify(v_name)[:15].upper() or "VENDOR"
                final_sap_id = f"SAP-AUTO-{clean_slug}-{uuid4().hex[:6].upper()}"
                op_status = OperationalStatus.ACTIVE
                mkt_status = MarketingStatus.PENDING
                appr_date = None
                type_label = "ACTIVE & PENDING (No Code, Generated ID)"

            # Check if vendor already exists by SAP ID or exact name
            existing = None
            if has_code:
                existing = Vendor.objects.filter(sap_vendor_id__iexact=final_sap_id).first()
            if not existing and v_name:
                v_name_clean = " ".join(v_name.split())
                existing = (
                    Vendor.objects.filter(vendor_name__iexact=v_name).first()
                    or Vendor.objects.filter(vendor_name__iexact=v_name_clean).first()
                )

            if existing:
                action = "UPDATE"
                updated_count += 1
                if not has_code:
                    final_sap_id = existing.sap_vendor_id
                if not dry_run:
                    existing.vendor_name = v_name
                    if has_code:
                        existing.sap_vendor_id = final_sap_id
                        existing.operational_status = OperationalStatus.ACTIVE
                        existing.marketing_status = MarketingStatus.APPROVED
                        existing.approved_at = now
                        if existing.onboarding_submission:
                            existing.onboarding_submission.finance_vendor_code = final_sap_id
                            existing.onboarding_submission.status = "activated"
                            existing.onboarding_submission.save(update_fields=["finance_vendor_code", "status"])
                    else:
                        existing.operational_status = OperationalStatus.ACTIVE
                        existing.marketing_status = MarketingStatus.PENDING
                    if v_email:
                        existing.email = v_email
                    if v_phone:
                        existing.phone = v_phone
                    if v_pan:
                        existing.pan = v_pan
                    if v_gstin:
                        existing.gstin = v_gstin
                    existing.save()
            else:
                action = "CREATE"
                if has_code:
                    created_approved += 1
                else:
                    created_pending += 1

                if not dry_run:
                    Vendor.objects.create(
                        org=org,
                        scope_node=scope_node,
                        vendor_name=v_name,
                        sap_vendor_id=final_sap_id,
                        email=v_email,
                        phone=v_phone,
                        pan=v_pan,
                        gstin=v_gstin,
                        operational_status=op_status,
                        marketing_status=mkt_status,
                        approved_at=appr_date,
                        address_line1=row_data.get("address", ""),
                        city=row_data.get("city", ""),
                        state=row_data.get("state", ""),
                        pincode=row_data.get("pincode", ""),
                        bank_name=row_data.get("bank_name", ""),
                        account_number=row_data.get("account_number", ""),
                        ifsc=row_data.get("ifsc", ""),
                        beneficiary_name=row_data.get("beneficiary_name", ""),
                    )

            print(f"[{action}] Row {row_idx:3d}: {v_name[:28]:<30} | Code: {final_sap_id:<24} | {type_label}")

    if dry_run:
        process_rows()
    else:
        with transaction.atomic():
            process_rows()

    print("\n" + "=" * 80)
    print("IMPORT SUMMARY:")
    print(f"  - Total Vendor rows processed       : {total_rows}")
    print(f"  - Newly Created WITH Code (Approved): {created_approved}")
    print(f"  - Newly Created WITHOUT Code (Pending): {created_pending}")
    print(f"  - Existing Vendors Updated          : {updated_count}")
    print(f"  - Mode                              : {'DRY RUN (No data written)' if dry_run else 'COMMITTED TO DATABASE'}")
    print("=" * 80)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import vendors from Excel spreadsheet.")
    parser.add_argument("file", help="Path to Excel file (.xlsx)")
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Simulate the import and preview results without modifying the database.",
    )
    args = parser.parse_args()

    import_vendors(args.file, dry_run=args.dry_run)
